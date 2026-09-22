import base64
import io
import json
from email import policy
from email.parser import BytesParser

import httpx
import pytest
from conftest import send
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.config import Settings
from app.file_parsers import parse_calamari, parse_holidays
from app.mailer import TEMPLATES, DeliveryError, GmailTransport, drain, queue_mail
from app.models import ImportBatch, LeaveRequest, Outbox


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Detailed timesheet"
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_document_parsers_support_utf8_and_exact_imported_hours():
    holidays = parse_holidays(
        "Fecha;Nombre\n2026-10-12;Día de prueba\n2026-02-30;Inválido".encode(), "festivos.csv"
    )
    assert holidays["holidays"] == [{"date": "2026-10-12", "name": "Día de prueba"}]
    assert len(holidays["errors"]) == 1
    data = workbook_bytes(
        [
            ["Nombre", "Equipos", "E-mail", "Tipo", "19/10/2026", "20/10/2026", "Suma"],
            ["Prueba", "Marketing", "worker@example.com", "Vacaciones", "3h", "4h", "7h"],
        ]
    )
    periods = parse_calamari(data, "prueba.xlsx")["periods"]
    assert len(periods) == 1
    assert periods[0]["allocation"] == [
        {"date": "2026-10-19", "hours": "3.0000"},
        {"date": "2026-10-20", "hours": "4.0000"},
    ]
    assert periods[0]["days"] == 0.875


@pytest.mark.postgres
def test_holiday_import_is_atomic_owned_and_idempotent(clients, database):
    admin = clients("admin")
    data = "Fecha,Nombre\n2026-10-12,Día sintético\n2026-12-25,Otro festivo".encode()
    uploaded = send(
        admin, "POST", "/api/holidays/import", files={"file": ("festivos.csv", data, "text/csv")}
    )
    assert uploaded.status_code == 200, uploaded.text
    batch_id = uploaded.json()["batchId"]
    assert len(uploaded.json()["holidays"]) == 2
    invalid = send(
        admin, "POST", "/api/holidays/import/confirm", json={"batchId": batch_id, "selected": [100]}
    )
    assert invalid.status_code == 400
    forged = send(
        admin,
        "POST",
        "/api/holidays/import/confirm",
        json={"batchId": batch_id, "selected": [0], "holidays": []},
    )
    assert forged.status_code == 422
    result = send(
        admin,
        "POST",
        "/api/holidays/import/confirm",
        json={"batchId": batch_id, "selected": [0, 1]},
    )
    assert result.status_code == 200 and result.json()["imported"] == 2
    repeated = send(
        admin,
        "POST",
        "/api/holidays/import/confirm",
        json={"batchId": batch_id, "selected": [0, 1]},
    )
    assert repeated.status_code == 200 and repeated.json()["duplicate"]
    assert (
        send(
            clients("worker"),
            "POST",
            "/api/holidays/import",
            files={"file": ("f.csv", data, "text/csv")},
        ).status_code
        == 403
    )


@pytest.mark.postgres
def test_calamari_import_roundtrip_preserves_fractions_and_detects_duplicates(clients, database):
    admin = clients("admin")
    data = workbook_bytes(
        [
            ["Nombre", "Equipos", "E-mail", "Tipo", "19/10/2026", "20/10/2026", "Suma"],
            ["Prueba", "Marketing", "worker@example.com", "Vacaciones", "3h", "4h", "7h"],
        ]
    )
    upload = send(
        admin,
        "POST",
        "/api/reports/calamari-import",
        files={
            "file": (
                "historico.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload.status_code == 200, upload.text
    payload = {"batchId": upload.json()["batchId"], "selected": [0]}
    confirmed = send(admin, "POST", "/api/reports/calamari-import/confirm", json=payload)
    assert confirmed.status_code == 200 and confirmed.json()["imported"] == 1, confirmed.text
    assert send(admin, "POST", "/api/reports/calamari-import/confirm", json=payload).json()[
        "duplicate"
    ]
    exported = admin.get("/api/reports/timesheet-export?ym=2026-10&userId=test-worker")
    assert exported.status_code == 200
    periods = parse_calamari(exported.content, "export.xlsx")["periods"]
    assert periods[0]["days"] == 0.875
    assert periods[0]["allocation"][0]["hours"] == "3.0000"
    assert admin.get("/api/reports/timesheet-export?ym=2026-13").status_code == 400
    with database.sessions() as db:
        assert db.scalar(select(func.count()).select_from(LeaveRequest)) == 1


@pytest.mark.postgres
def test_bad_calamari_batch_rolls_back_all_rows(clients, database):
    admin = clients("admin")
    data = workbook_bytes(
        [
            ["Nombre", "Equipos", "E-mail", "Tipo", "19/10/2026", "Suma"],
            ["Prueba", "Marketing", "worker@example.com", "Vacaciones", "8h", "8h"],
            ["Desconocido", "Marketing", "missing@example.com", "Vacaciones", "8h", "8h"],
        ]
    )
    upload = send(
        admin,
        "POST",
        "/api/reports/calamari-import",
        files={"file": ("historico.xlsx", data, "application/octet-stream")},
    )
    assert upload.status_code == 200, upload.text
    confirmed = send(
        admin,
        "POST",
        "/api/reports/calamari-import/confirm",
        json={"batchId": upload.json()["batchId"], "selected": [0, 1]},
    )
    assert confirmed.status_code == 409
    with database.sessions() as db:
        assert db.scalar(select(func.count()).select_from(LeaveRequest)) == 0
        assert not db.get(ImportBatch, upload.json()["batchId"]).confirmed


@pytest.mark.postgres
def test_export_does_not_turn_employee_text_into_formulas(clients, database):
    from app.models import Employee

    admin = clients("admin")
    with database.sessions() as db:
        db.get(Employee, "test-worker").name = '=HYPERLINK("https://example.invalid","name")'
        db.commit()
    response = admin.get("/api/reports/timesheet-export?ym=2026-10&userId=test-worker")
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["A2"].value.startswith("=HYPERLINK")
    workbook.close()


def gmail_settings():
    return Settings(
        database_url="",
        environment="test",
        app_url="http://testserver",
        mail_enabled=True,
        gmail_from="notifications@example.com",
        google_client_id="synthetic-client",
        google_client_secret="synthetic-secret-for-tests",
        gmail_refresh_token="synthetic-refresh-for-tests",
    )


def test_gmail_api_uses_oauth_and_rfc_mime_without_leaking_private_notes():
    calls = []

    def handle(request):
        calls.append(request)
        if request.url.host == "oauth2.googleapis.com":
            assert b"grant_type=refresh_token" in request.content
            return httpx.Response(
                200, json={"access_token": "synthetic-access", "expires_in": 3600}
            )
        assert request.url.host == "gmail.googleapis.com"
        assert request.headers["authorization"] == "Bearer synthetic-access"
        raw = json.loads(request.content)["raw"]
        message = BytesParser(policy=policy.default).parsebytes(base64.urlsafe_b64decode(raw))
        assert message["To"] == "worker@example.com"
        html = message.get_body(preferencelist=("html",)).get_content()
        assert "&lt;script&gt;" in html and "<script>" not in html
        return httpx.Response(200, json={"id": "synthetic-message"})

    client = httpx.Client(transport=httpx.MockTransport(handle))
    transport = GmailTransport(gmail_settings(), client)
    html = TEMPLATES.get_template("notification.html").render(
        name="<script>",
        worker_name="Prueba",
        event="created",
        date_from="2026-10-19",
        date_to="2026-10-20",
        app_url="http://testserver",
    )
    assert (
        transport.send("worker@example.com", "Prueba", html, "Prueba", "test-message")
        == "synthetic-message"
    )
    assert len(calls) == 2
    client.close()


@pytest.mark.parametrize(
    "status,state", [(429, "pending"), (400, "failed"), (403, "failed"), (500, "uncertain")]
)
def test_gmail_classifies_safe_retries_and_uncertain_delivery(status, state):
    def handle(request):
        if request.url.host == "oauth2.googleapis.com":
            return httpx.Response(200, json={"access_token": "synthetic", "expires_in": 3600})
        return httpx.Response(status, json={"error": {"errors": []}})

    with httpx.Client(transport=httpx.MockTransport(handle)) as client:
        transport = GmailTransport(gmail_settings(), client)
        with pytest.raises(DeliveryError) as caught:
            transport.send("worker@example.com", "Test", "Test", "Test", "id")
        assert caught.value.state == state


@pytest.mark.postgres
def test_outbox_delivery_is_bounded_and_does_not_retry_uncertain_messages(database, employees):
    class FakeTransport:
        def __init__(self):
            self.calls = 0

        def send(self, *args):
            self.calls += 1
            raise DeliveryError("uncertain", "GMAIL_DELIVERY_UNCERTAIN")

        def close(self):
            pass

    with database.sessions() as db:
        queue_mail(
            db,
            "synthetic-event",
            "worker@example.com",
            "Aviso",
            {
                "name": "Prueba",
                "worker_name": "Prueba",
                "event": "created",
                "date_from": "2026-10-19",
                "date_to": "2026-10-20",
            },
        )
        db.commit()
    fake = FakeTransport()
    result = drain(database, gmail_settings(), transport=fake)
    assert result["uncertain"] == 1 and fake.calls == 1
    result = drain(database, gmail_settings(), transport=fake)
    assert result["uncertain"] == 0 and fake.calls == 1
    with database.sessions() as db:
        item = db.scalar(select(Outbox))
        assert item.status == "uncertain" and item.attempts == 1
