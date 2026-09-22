"""Bounded document parsers, executed only in a supervised resource-limited process."""

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal

MAX_ROWS = 5000
MAX_COLUMNS = 400


def normalize(value):
    return "".join(
        c
        for c in unicodedata.normalize("NFKD", str(value or "").lower())
        if not unicodedata.combining(c)
    ).strip()


def parse_date(value):
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        text = str(value or "").strip()
        result = None
        for format_string in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
            try:
                result = datetime.strptime(text, format_string).date()
                break
            except ValueError:
                continue
        if result is None:
            raise ValueError("Fecha no reconocida")
    if not 1900 <= result.year <= 2200:
        raise ValueError("Año fuera del intervalo admitido")
    return result.isoformat()


def safe_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        total = sum(e.file_size for e in entries)
        if len(entries) > 1000 or total > 64 * 1024 * 1024:
            raise ValueError("El documento comprimido es demasiado grande")
        if any(
            e.flag_bits & 1
            or e.file_size > 32 * 1024 * 1024
            or e.file_size > max(e.compress_size, 1) * 500
            for e in entries
        ):
            raise ValueError("Documento comprimido no admitido")


def spreadsheet_rows(data, name):
    suffix = name.lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = []
        for index, row in enumerate(reader):
            if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
                raise ValueError("Demasiadas filas o columnas")
            rows.append(row)
        return rows
    if suffix == "xlsx":
        safe_zip(data)
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = next(
                (s for s in workbook if "num" in normalize(s.title)), workbook.worksheets[0]
            )
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Demasiadas filas o columnas")
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise ValueError("Demasiadas filas o columnas")
                rows.append(list(row))
            return rows
        finally:
            workbook.close()
    if suffix == "xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                raise ValueError("Demasiadas filas o columnas")
            rows = []
            for row_index in range(sheet.nrows):
                row = []
                for cell in sheet.row(row_index):
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    row.append(value)
                rows.append(row)
            return rows
        finally:
            workbook.release_resources()
    raise ValueError("Usa XLSX, XLS o CSV")


def pdf_reader(data):
    if not data.startswith(b"%PDF-"):
        raise ValueError("El archivo no contiene un PDF válido")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted or not 1 <= len(reader.pages) <= 50:
        raise ValueError("PDF cifrado o con demasiadas páginas")
    root = reader.trailer["/Root"]
    if "/OpenAction" in root or "/AA" in root:
        raise ValueError("No se admiten acciones activas en PDF")
    names = root.get("/Names", {})
    if hasattr(names, "get_object"):
        names = names.get_object()
    if "/JavaScript" in names or "/EmbeddedFiles" in names:
        raise ValueError("No se admiten scripts ni archivos incrustados en PDF")
    for page in reader.pages:
        if "/AA" in page:
            raise ValueError("No se admiten acciones activas en PDF")
    return reader


def parse_holidays(data, name):
    if name.lower().endswith(".pdf"):
        reader = pdf_reader(data)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if len(text) > 1_000_000:
            raise ValueError("El PDF contiene demasiado texto")
        rows = []
        for line in text.splitlines():
            match = re.search(r"\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[/.-]\d{1,2}[/.-]\d{4}", line)
            if match:
                rows.append([match.group(), line.replace(match.group(), "", 1).strip(" :-–—\t")])
    else:
        rows = spreadsheet_rows(data, name)
    results, errors = [], []
    seen = set()
    for index, row in enumerate(rows):
        if not row or not any(str(v or "").strip() for v in row):
            continue
        if index == 0 and normalize(row[0]) in {"fecha", "date", "dia"}:
            continue
        try:
            day = parse_date(row[0])
            label = str(row[1] or "").strip()
            if not 1 <= len(label) <= 200:
                raise ValueError("Nombre de festivo no válido")
            if day not in seen:
                results.append({"date": day, "name": label})
                seen.add(day)
        except (ValueError, IndexError):
            errors.append({"row": index + 1, "error": "Fecha o nombre no válido"})
    if not results:
        raise ValueError("No se han reconocido festivos válidos")
    return {"holidays": sorted(results, key=lambda r: r["date"]), "errors": errors}


def duration_hours(value):
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        value = str(value or "").strip().lower().replace(",", ".")
        if not value:
            return Decimal(0)
        if re.fullmatch(r"\d+(?:\.\d+)?", value):
            amount = Decimal(value)
        else:
            match = re.fullmatch(r"\s*(?:(\d+(?:\.\d+)?)h)?\s*(?:(\d+(?:\.\d+)?)m)?\s*", value)
            if not match or not any(match.groups()):
                raise ValueError("Duración no válida")
            amount = Decimal(match[1] or 0) + Decimal(match[2] or 0) / 60
    if not amount.is_finite() or not 0 <= amount <= 24:
        raise ValueError("Horas fuera del intervalo admitido")
    return amount.quantize(Decimal("0.0001"))


def parse_calamari(data, name):
    rows = spreadsheet_rows(data, name)
    if len(rows) < 2:
        raise ValueError("La hoja no contiene datos")
    header = [normalize(v) for v in rows[0]]
    email_column = next(
        (
            i
            for i, v in enumerate(header)
            if v in {"e-mail", "email", "correo", "correo electronico"}
        ),
        None,
    )
    if email_column is None:
        raise ValueError("Falta la columna de correo")
    dates = []
    for i, value in enumerate(rows[0]):
        try:
            dates.append((i, parse_date(value)))
        except ValueError:
            continue
    if not dates or len(dates) > 367:
        raise ValueError("Faltan las columnas de fechas o hay demasiadas")
    date_values = [d for _, d in dates]
    if date_values != sorted(set(date_values)):
        raise ValueError("Las fechas deben estar ordenadas y no repetirse")
    label_column = min(i for i, _ in dates) - 1
    if label_column <= email_column:
        raise ValueError("No se reconoce la columna del tipo de ausencia")
    name_column = next(
        (i for i, v in enumerate(header) if v in {"nombre", "name", "first name"}), 0
    )
    surname_column = next(
        (i for i, v in enumerate(header) if v in {"apellidos", "surname", "last name"}), None
    )
    periods, errors = [], []
    for index, row in enumerate(rows[1:], 2):
        if len(row) <= label_column:
            continue
        label = normalize(row[label_column])
        absence_type = next(
            (
                kind
                for kind, word in [
                    ("vacaciones", "vacacion"),
                    ("baja", "baja"),
                    ("ausencia", "ausencia"),
                    ("permiso", "permiso"),
                ]
                if word in label
            ),
            None,
        )
        if not absence_type:
            continue
        email = str(row[email_column] or "").strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            errors.append({"row": index, "error": "Correo no válido"})
            continue
        employee_name = str(row[name_column] or "").strip()
        if surname_column is not None:
            employee_name += " " + str(row[surname_column] or "").strip()
        allocation = []
        invalid = False
        for column, day in dates:
            try:
                hours = duration_hours(row[column] if column < len(row) else 0)
                if hours:
                    allocation.append({"date": day, "hours": str(hours)})
            except ValueError:
                errors.append({"row": index, "error": "Duración no válida"})
                invalid = True
                break
        if invalid or not allocation:
            continue
        # Split at non-weekend gaps; keep exact hours for every imported day.
        groups = []
        for day in allocation:
            current_date = date.fromisoformat(day["date"])
            if not groups:
                groups.append([])
            elif groups[-1]:
                previous = date.fromisoformat(groups[-1][-1]["date"])
                gap = [
                    previous + timedelta(days=i) for i in range(1, (current_date - previous).days)
                ]
                if any(d.weekday() < 5 for d in gap):
                    groups.append([])
            groups[-1].append(day)
        for group in groups:
            total = sum((Decimal(d["hours"]) for d in group), Decimal(0))
            periods.append(
                {
                    "email": email,
                    "name": employee_name.strip(),
                    "type": absence_type,
                    "dateFrom": group[0]["date"],
                    "dateTo": group[-1]["date"],
                    "days": float(total / 8),
                    "halfStart": Decimal(group[0]["hours"]) == 4,
                    "halfEnd": len(group) > 1 and Decimal(group[-1]["hours"]) == 4,
                    "allocation": group,
                    "sourceRow": index,
                }
            )
    if not periods:
        raise ValueError("No se han reconocido ausencias válidas")
    return {"periods": periods, "errors": errors}


def inspect_attachment(data):
    if data.startswith(b"%PDF-"):
        pdf_reader(data)
        return {"mime": "application/pdf", "extension": ".pdf"}
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = 16_000_000
    with Image.open(io.BytesIO(data)) as image:
        formats = {
            "JPEG": ("image/jpeg", ".jpg"),
            "PNG": ("image/png", ".png"),
            "WEBP": ("image/webp", ".webp"),
            "GIF": ("image/gif", ".gif"),
        }
        if image.format not in formats or image.width * image.height > 16_000_000:
            raise ValueError("Imagen no admitida")
        mime, extension = formats[image.format]
        image.verify()
    return {"mime": mime, "extension": extension}
